"""Parse the annotation Q&A xlsx into structured records with associated images.

The xlsx stores Q&A rows across several sheets and embeds example images that are
anchored to specific cells (drawingN.xml). We map every embedded image back to the
row it is anchored to, then bundle each row's text + images into one record.
"""
from __future__ import annotations

import json
import re
import shutil
import zipfile
from dataclasses import dataclass, field, asdict
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.utils import get_column_letter

NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}


@dataclass
class Record:
    source: str            # "xlsx:<sheet>"
    sheet: str
    row: int               # 1-based excel row
    fields: dict           # column-letter -> text
    text: str              # concatenated readable text
    images: list = field(default_factory=list)  # relative image paths


def _rels_map(zf: zipfile.ZipFile, rels_path: str) -> dict:
    """rId -> target (image media path or external url)."""
    out = {}
    if rels_path not in zf.namelist():
        return out
    root = ET.fromstring(zf.read(rels_path))
    for rel in root:
        out[rel.get("Id")] = rel.get("Target")
    return out


def _sheet_to_drawing(zf: zipfile.ZipFile, sheet_file: str) -> str | None:
    rels = f"xl/worksheets/_rels/{Path(sheet_file).name}.rels"
    m = _rels_map(zf, rels)
    for v in m.values():
        if "drawings/drawing" in v:
            return v.replace("../", "xl/")
    return None


def _parse_drawing(zf: zipfile.ZipFile, drawing_path: str) -> list[tuple[int, str]]:
    """Return list of (anchor_row_0based, media_path) for each picture."""
    drawing_rels = f"xl/drawings/_rels/{Path(drawing_path).name}.rels"
    rid_to_target = _rels_map(zf, drawing_rels)
    root = ET.fromstring(zf.read(drawing_path))
    results = []
    for anchor in root:
        frm = anchor.find("xdr:from", NS)
        if frm is None:
            continue
        row_el = frm.find("xdr:row", NS)
        if row_el is None:
            continue
        row0 = int(row_el.text)
        blip = anchor.find(".//a:blip", NS)
        if blip is None:
            continue
        embed = blip.get(f"{{{NS['r']}}}embed")
        target = rid_to_target.get(embed)
        if not target:
            continue
        media = target.replace("../", "xl/")
        results.append((row0, media))
    return results


def _cell_text(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return s


def ingest(xlsx_path: str, out_dir: str) -> list[Record]:
    xlsx_path = Path(xlsx_path)
    out = Path(out_dir)
    img_dir = out / "images"
    img_dir.mkdir(parents=True, exist_ok=True)

    zf = zipfile.ZipFile(xlsx_path)

    # map workbook sheet name -> sheetN.xml
    wb_rels = _rels_map(zf, "xl/_rels/workbook.xml.rels")
    wb_xml = zf.read("xl/workbook.xml").decode("utf-8")
    name_to_rid = re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wb_xml)
    sheetname_to_file = {}
    for name, rid in name_to_rid:
        tgt = wb_rels.get(rid, "")
        if tgt:
            sheetname_to_file[name] = "xl/" + tgt if not tgt.startswith("xl/") else tgt

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    records: list[Record] = []
    img_counter = 0

    for ws in wb.worksheets:
        sheet_name = ws.title
        sheet_file = sheetname_to_file.get(sheet_name)
        # row -> list of media paths
        row_images: dict[int, list[str]] = {}
        if sheet_file:
            drawing = _sheet_to_drawing(zf, sheet_file)
            if drawing:
                for row0, media in _parse_drawing(zf, drawing):
                    row_images.setdefault(row0 + 1, []).append(media)

        # extract & copy images, build saved path map per row
        row_saved: dict[int, list[str]] = {}
        for row1, medias in row_images.items():
            for media in medias:
                img_counter += 1
                ext = Path(media).suffix or ".jpeg"
                safe_sheet = re.sub(r"[^\w]", "_", sheet_name)
                fname = f"{safe_sheet}_r{row1}_{img_counter}{ext}"
                dest = img_dir / fname
                try:
                    with zf.open(media) as src, open(dest, "wb") as dst:
                        shutil.copyfileobj(src, dst)
                    row_saved.setdefault(row1, []).append(f"images/{fname}")
                except KeyError:
                    pass

        for r in range(1, ws.max_row + 1):
            fields = {}
            for c in range(1, ws.max_column + 1):
                t = _cell_text(ws.cell(r, c).value)
                if t:
                    fields[get_column_letter(c)] = t
            imgs = row_saved.get(r, [])
            if not fields and not imgs:
                continue
            text = "\n".join(f"{k}: {v}" for k, v in fields.items())
            records.append(
                Record(
                    source=f"xlsx:{sheet_name}",
                    sheet=sheet_name,
                    row=r,
                    fields=fields,
                    text=text,
                    images=imgs,
                )
            )

    out.mkdir(parents=True, exist_ok=True)
    with open(out / "xlsx_records.jsonl", "w", encoding="utf-8") as f:
        for rec in records:
            f.write(json.dumps(asdict(rec), ensure_ascii=False) + "\n")

    n_img = sum(len(r.images) for r in records)
    print(f"sheets={len(wb.worksheets)} records={len(records)} images_linked={n_img}")
    return records


if __name__ == "__main__":
    import sys
    import config
    xlsx = sys.argv[1] if len(sys.argv) > 1 else config.XLSX_PATH
    out = sys.argv[2] if len(sys.argv) > 2 else str(config.KB_DIR)
    ingest(xlsx, out)
